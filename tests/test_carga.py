"""Carga de CSV/XLSX (`app/carga.py`): lectura como función pura (AppTest no puede subir archivos;
la subida real se prueba con Playwright), todos los rechazos y la seguridad del XLSX. Los archivos
se generan aquí, en `tmp_path` o en memoria: no hay binarios en el repo."""

import io
import re
import sys
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

pytest.importorskip("streamlit")
openpyxl = pytest.importorskip("openpyxl")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "app"))

import bond_datos as bd  # noqa: E402
import i18n  # noqa: E402
from carga import (  # noqa: E402
    MAX_BYTES,
    MAX_FILAS,
    Columna,
    ErrorCarga,
    leer_tabla,
    mensaje_error,
    plantilla_csv,
)
from muestreo_datos import COLUMNAS_ARCHIVO, COLUMNAS_TABLA, tabla_de_levantamiento  # noqa: E402

from conminucion.plantas_reales import VAZANTE_C1  # noqa: E402

PEGADA = tabla_de_levantamiento(VAZANTE_C1)  # «la tabla pegada a mano»
FINOS = bd.tabla_inicial()


def csv_bytes(df: pd.DataFrame, cols, sep: str, coma: bool, lang="es", bom=False, enc="utf-8"):
    filas = [sep.join(c.titulo(lang) for c in cols)]
    for _, r in df.iterrows():
        celdas = ["" if pd.isna(r[c.clave]) else f"{r[c.clave]:g}" for c in cols]
        filas.append(sep.join(x.replace(".", ",") if coma else x for x in celdas))
    texto = ("﻿" if bom else "") + "\r\n".join(filas) + "\r\n"
    return texto.encode(enc)


def xlsx_bytes(df: pd.DataFrame, cols, lang="es") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([c.titulo(lang) for c in cols])
    for _, r in df.iterrows():
        ws.append([None if pd.isna(r[c.clave]) else float(r[c.clave]) for c in cols])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def igual(a: pd.DataFrame, b: pd.DataFrame) -> None:
    pd.testing.assert_frame_equal(a.reset_index(drop=True), b.reset_index(drop=True))


# ------------------------------------------------------------------ lo que se lee
@pytest.mark.parametrize(
    ("sep", "coma", "lang"),
    [(";", True, "es"), (",", False, "en"), (";", True, "en"), (",", False, "es")],
)
@pytest.mark.parametrize("bom", [False, True])
def test_csv_en_los_dos_formatos_de_excel_da_la_tabla_pegada(sep, coma, lang, bom):
    datos = csv_bytes(PEGADA, COLUMNAS_TABLA, sep, coma, lang, bom)
    igual(leer_tabla("muestreo.csv", datos, COLUMNAS_TABLA), PEGADA)


@pytest.mark.parametrize("lang", ["es", "en"])
def test_xlsx_da_la_misma_tabla_que_el_csv_y_la_pegada(lang):
    datos = xlsx_bytes(PEGADA, COLUMNAS_TABLA, lang)
    leida = leer_tabla("muestreo.XLSX", datos, COLUMNAS_TABLA)
    igual(leida, PEGADA)
    igual(leida, leer_tabla("m.csv", csv_bytes(PEGADA, COLUMNAS_TABLA, ";", True), COLUMNAS_TABLA))


def test_csv_latin1_y_una_sola_columna_con_coma_decimal():
    col = (Columna("d", "Tamaño (µm)", "Size (µm)"),)
    datos = b"Tama\xf1o (\xb5m)\r\n583,9\r\n"  # Latin-1: no es UTF-8 válido
    with pytest.raises(UnicodeDecodeError):
        datos.decode("utf-8")
    assert leer_tabla("a.csv", datos, col)["d"].tolist() == [583.9]
    sin_sep = b"finos\r\n583,9\r\n464,5\r\n"
    esperado = pd.DataFrame({"finos": [583.9, 464.5]})
    igual(leer_tabla("a.csv", sin_sep, bd.COLUMNAS_CICLOS), esperado)
    igual(leer_tabla("a.csv", b"finos\n583.9\n464.5\n", bd.COLUMNAS_CICLOS), esperado)


def test_bond_csv_y_xlsx_dan_la_tabla_de_ciclos():
    cols = bd.COLUMNAS_CICLOS
    igual(leer_tabla("b.csv", csv_bytes(FINOS, cols, ";", True), cols), FINOS)
    igual(leer_tabla("b.xlsx", xlsx_bytes(FINOS, cols, "en"), cols), FINOS)


def test_las_columnas_opcionales_ausentes_quedan_en_nan_y_el_orden_del_archivo_no_importa():
    cols = COLUMNAS_TABLA[:4]
    solo = PEGADA[[c.clave for c in cols]]
    datos = csv_bytes(solo[solo.columns[::-1]], cols[::-1], ";", True)
    leida = leer_tabla("x.csv", datos, COLUMNAS_TABLA)
    igual(leida, PEGADA)
    assert leida[["molino_alim", "molino_desc"]].isna().all().all()


def test_las_filas_totalmente_vacias_se_saltan_y_el_error_cuenta_la_fila_de_la_hoja():
    datos = b"finos\n1\n\n2\nx\n"
    with pytest.raises(ErrorCarga) as e:
        leer_tabla("a.csv", datos, bd.COLUMNAS_CICLOS)
    assert e.value.campos["fila"] == 5


def test_la_plantilla_sale_de_las_mismas_columnas_y_se_lee_de_vuelta():
    for lang, sep in (("es", ";"), ("en", ",")):
        datos = plantilla_csv(COLUMNAS_ARCHIVO, PEGADA, lang)
        assert datos.startswith(b"\xef\xbb\xbf")
        assert datos.decode("utf-8-sig").splitlines()[0].split(sep)[0] == COLUMNAS_ARCHIVO[
            0
        ].titulo(lang)
        igual(leer_tabla("plantilla.csv", datos, COLUMNAS_TABLA), PEGADA)


# ------------------------------------------------------------------ los rechazos
@pytest.mark.parametrize("nombre", ["a.xlsm", "a.xls", "a.txt", "a", "a.csv.exe", "a.CSV.zip"])
def test_rechaza_otras_extensiones(nombre):
    with pytest.raises(ErrorCarga) as e:
        leer_tabla(nombre, b"finos\n1\n", bd.COLUMNAS_CICLOS)
    assert e.value.clave == "carga.err.ext"


def test_rechaza_mas_de_1_mb_en_el_codigo_y_acepta_justo_el_limite():
    base = b"finos\n1\n"
    justo = base + b"\n" * (MAX_BYTES - len(base))
    assert len(leer_tabla("a.csv", justo, bd.COLUMNAS_CICLOS)) == 1
    with pytest.raises(ErrorCarga) as e:
        leer_tabla("a.csv", justo + b"\n", bd.COLUMNAS_CICLOS)
    assert e.value.clave == "carga.err.tamano"


def test_rechaza_mas_de_500_filas_en_csv_y_xlsx():
    ok = pd.DataFrame({"finos": np.arange(MAX_FILAS, dtype=float)})
    mal = pd.DataFrame({"finos": np.arange(MAX_FILAS + 1, dtype=float)})
    cols = bd.COLUMNAS_CICLOS
    assert len(leer_tabla("a.csv", csv_bytes(ok, cols, ";", True), cols)) == MAX_FILAS
    assert len(leer_tabla("a.xlsx", xlsx_bytes(ok, cols), cols)) == MAX_FILAS
    for datos, nombre in (
        (csv_bytes(mal, cols, ";", True), "a.csv"),
        (xlsx_bytes(mal, cols), "a.xlsx"),
    ):
        with pytest.raises(ErrorCarga) as e:
            leer_tabla(nombre, datos, cols)
        assert e.value.clave == "carga.err.filas" and e.value.campos["max"] == MAX_FILAS


def test_rechaza_columnas_obligatorias_faltantes_y_las_nombra():
    datos = csv_bytes(PEGADA[["tamiz", "alim"]], COLUMNAS_TABLA[:2], ";", True)
    with pytest.raises(ErrorCarga) as e:
        leer_tabla("a.csv", datos, COLUMNAS_TABLA)
    assert e.value.clave == "carga.err.faltan"
    assert [c.clave for c in e.value.campos["cols"]] == ["uf", "of"]
    with i18n.use("en"):
        assert "Underflow (% passing), Overflow (% passing)" in mensaje_error(e.value)
    with i18n.use("es"):
        assert "Underflow (% pasante), Overflow (% pasante)" in mensaje_error(e.value)


@pytest.mark.parametrize(
    "celda", ["abc", "", "NaN", "inf", "1e999", "1_0", "12/10/2020", "--1", "0x1F"]
)
def test_rechaza_celdas_no_numericas_con_fila_y_columna(celda):
    datos = f"tamiz;alim;uf;of\n100;100;100;100\n50;{celda};60;70\n".encode()
    with pytest.raises(ErrorCarga) as e:
        leer_tabla("a.csv", datos, COLUMNAS_TABLA)
    assert e.value.clave == "carga.err.celda"
    assert e.value.campos["fila"] == 3 and e.value.campos["col"].clave == "alim"
    with i18n.use("es"):
        assert "La fila 3" in mensaje_error(
            e.value
        ) and "Alimentación (% pasante)" in mensaje_error(e.value)
    with i18n.use("en"):
        assert "Row 3" in mensaje_error(e.value) and "Feed (% passing)" in mensaje_error(e.value)


def test_rechaza_celdas_no_numericas_en_xlsx_texto_fecha_y_booleano():
    import datetime

    for valor in ("texto", datetime.date(2020, 1, 2), True):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["finos"])
        ws.append([1.5])
        ws.append([valor])
        out = io.BytesIO()
        wb.save(out)
        with pytest.raises(ErrorCarga) as e:
            leer_tabla("a.xlsx", out.getvalue(), bd.COLUMNAS_CICLOS)
        assert e.value.clave == "carga.err.celda" and e.value.campos["fila"] == 3


def test_rechaza_archivos_vacios_ilegibles_y_zip_bombs():
    for nombre, datos, clave in (
        ("a.csv", b"", "carga.err.vacio"),
        ("a.csv", b"finos\n", "carga.err.vacio"),
        ("a.xlsx", b"esto no es un zip", "carga.err.ilegible"),
        ("a.xlsx", b"PK\x03\x04roto", "carga.err.ilegible"),
    ):
        with pytest.raises(ErrorCarga) as e:
            leer_tabla(nombre, datos, bd.COLUMNAS_CICLOS)
        assert e.value.clave == clave
    bomba = io.BytesIO()
    with zipfile.ZipFile(bomba, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("relleno.bin", b"\0" * (60 * 1024 * 1024))
    assert len(bomba.getvalue()) < MAX_BYTES
    with pytest.raises(ErrorCarga) as e:
        leer_tabla("a.xlsx", bomba.getvalue(), bd.COLUMNAS_CICLOS)
    assert e.value.clave == "carga.err.tamano"


def test_todos_los_rechazos_tienen_texto_en_los_dos_idiomas_y_sin_traza():
    claves = {
        "carga.err." + k
        for k in ("ext", "tamano", "filas", "faltan", "celda", "ilegible", "vacio", "xml")
    }
    for lang in ("es", "en"):
        for k in claves:
            texto = i18n.TEXTS[lang][k]
            assert "Traceback" not in texto and "Error" not in texto.split(":")[0]


# ------------------------------------------------------------------ seguridad del XLSX
def test_openpyxl_usa_defusedxml():
    import openpyxl.xml

    assert openpyxl.xml.DEFUSEDXML is True


def test_si_defusedxml_no_esta_activo_el_xlsx_se_rechaza(monkeypatch):
    import openpyxl.xml

    monkeypatch.setattr(openpyxl.xml, "DEFUSEDXML", False)
    with pytest.raises(ErrorCarga) as e:
        leer_tabla("a.xlsx", xlsx_bytes(FINOS, bd.COLUMNAS_CICLOS), bd.COLUMNAS_CICLOS)
    assert e.value.clave == "carga.err.xml"


def test_una_formula_devuelve_el_valor_guardado_y_no_la_formula(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["finos"])
    ws.append(["=40+2"])
    ws.append([7.5])
    ruta = tmp_path / "f.xlsx"
    wb.save(ruta)
    # openpyxl no guarda el valor calculado de una fórmula (Excel sí): se agrega a mano
    with zipfile.ZipFile(ruta) as z:
        partes = {n: z.read(n) for n in z.namelist()}
    hoja = partes["xl/worksheets/sheet1.xml"].decode()
    assert "<f>40+2</f><v />" in hoja
    partes["xl/worksheets/sheet1.xml"] = hoja.replace(
        "<f>40+2</f><v />", "<f>40+2</f><v>99</v>"
    ).encode()
    with zipfile.ZipFile(ruta, "w") as z:
        for n, b in partes.items():
            z.writestr(n, b)
    leida = leer_tabla("f.xlsx", ruta.read_bytes(), bd.COLUMNAS_CICLOS)
    assert leida["finos"].tolist() == [99.0, 7.5]  # el valor guardado, no 42 ni el texto


def test_una_formula_sin_valor_guardado_no_se_evalua_y_su_fila_vacia_se_salta():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["finos"])
    ws.append(["=1+1"])
    ws.append([7.5])
    out = io.BytesIO()
    wb.save(out)
    leida = leer_tabla("f.xlsx", out.getvalue(), bd.COLUMNAS_CICLOS)
    assert leida["finos"].tolist() == [7.5]  # nunca 2: no se calcula la fórmula


def test_el_xlsx_con_una_entidad_externa_no_se_resuelve(tmp_path):
    xxe = '<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x SYSTEM "file:///etc/hostname">]>'
    ruta = tmp_path / "x.xlsx"
    ruta.write_bytes(xlsx_bytes(FINOS, bd.COLUMNAS_CICLOS))
    with zipfile.ZipFile(ruta) as z:
        partes = {n: z.read(n) for n in z.namelist()}
    hoja = partes["xl/worksheets/sheet1.xml"].decode()
    partes["xl/worksheets/sheet1.xml"] = (
        re.sub(r"^<\?xml[^>]*\?>", xxe, hoja).replace("<v>", "<v>&x;", 1).encode()
    )
    with zipfile.ZipFile(ruta, "w") as z:
        for n, b in partes.items():
            z.writestr(n, b)
    with pytest.raises(ErrorCarga) as e:
        leer_tabla("x.xlsx", ruta.read_bytes(), bd.COLUMNAS_CICLOS)
    assert e.value.clave in ("carga.err.ilegible", "carga.err.celda")


def test_columna_alias_ignora_mayusculas_espacios_y_bom():
    c = Columna("x", "Tamiz (µm)", "Sieve (µm)")
    # casefold unifica el micro (U+00B5) y la mu griega (U+03BC): ambos encabezados valen
    assert all(a in c.alias for a in ("x", "tamiz (μm)", "sieve (μm)"))
    datos = "﻿  TAMIZ   (µm) ;Sieve (µm)\n1;2\n".encode()
    cols = (c, Columna("y", "Otra", "Other", False))
    assert leer_tabla("a.csv", datos, cols)["x"].tolist() == [1.0]
