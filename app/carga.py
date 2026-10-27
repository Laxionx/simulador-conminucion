"""Carga de tablas desde CSV o XLSX: `leer_tabla` (función pura, sin Streamlit) y el cargador de la
interfaz (aviso de privacidad, `st.file_uploader` y la plantilla descargable).

Reglas de seguridad: solo `.csv` y `.xlsx` (nunca `.xlsm` ni `.xls`); a lo más 1 MB y 500 filas; el
XLSX se lee con `openpyxl` en modo de solo lectura, con `defusedxml` activo (si no lo está, se
rechaza) y sin evaluar fórmulas (se usa el valor que el archivo guardó); no se descomprime un XLSX
cuyo contenido pase de `MAX_DESCOMPRIMIDO`. Lo subido vive solo en memoria: no se guarda en disco ni
se dibuja como HTML (el nombre del archivo ni siquiera se muestra). Los errores son una
`ErrorCarga` con la clave de un texto traducido y sus campos; nunca una traza.
"""

from __future__ import annotations

import csv
import io
import math
import re
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import PurePath

import numpy as np
import pandas as pd

MAX_BYTES = 1024 * 1024  # 1 MB, igual que `maxUploadSize` de .streamlit/config.toml
MAX_FILAS = 500
MAX_COLUMNAS = 40  # se leen a lo más estas columnas de una hoja
MAX_DESCOMPRIMIDO = 50 * 1024 * 1024  # tope del contenido de un XLSX descomprimido (zip bomb)
EXTENSIONES = (".csv", ".xlsx")
_NUMERO = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?")


@dataclass(frozen=True)
class Columna:
    """Una columna de una tabla: su clave interna y sus encabezados en español e inglés. Se
    aceptan los tres al leer; la plantilla usa el del idioma de la app."""

    clave: str
    es: str
    en: str
    obligatoria: bool = True

    def titulo(self, lang: str) -> str:
        return self.en if lang == "en" else self.es

    @property
    def alias(self) -> frozenset[str]:
        return frozenset(_normal(x) for x in (self.clave, self.es, self.en))


@dataclass
class ErrorCarga(ValueError):
    """Un rechazo de la carga: `clave` es el texto traducido (`carga.err.*`) y `campos` rellenan
    sus {campos}. Un valor `Columna` (o una lista de ellas) se escribe con su encabezado."""

    clave: str
    campos: dict[str, object] = field(default_factory=dict)

    def __str__(self) -> str:
        return self.clave


def _normal(x: str) -> str:
    return re.sub(r"\s+", " ", x.replace("﻿", "")).strip().casefold()


def _numero(x: object, coma: bool) -> float | None:
    """El número de una celda, o None si no lo es (texto, vacío, NaN, infinito, booleano)."""
    if isinstance(x, bool):
        return None
    if isinstance(x, int | float | np.integer | np.floating):
        return float(x) if math.isfinite(float(x)) else None
    if not isinstance(x, str):
        return None
    s = x.strip().replace("\xa0", "").replace(" ", "")
    if coma and s.count(",") == 1 and "." not in s:
        s = s.replace(",", ".")
    return float(s) if _NUMERO.fullmatch(s) and math.isfinite(float(s)) else None


def _vacia(x: object) -> bool:
    return x is None or (isinstance(x, str) and not x.strip())


# ------------------------------------------------------------------ lectura de bytes
def _filas_csv(contenido: bytes, n_cols: int) -> tuple[list[list[object]], bool]:
    """(filas, coma_decimal). Separador «;» (decimal coma, Excel en español) o «,» (decimal punto,
    Excel en inglés); con una sola columna esperada y sin separador, cada línea es una celda."""
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = contenido.decode("latin-1")
    cabecera = texto.lstrip("\r\n").split("\n", 1)[0]
    if ";" in cabecera:
        sep, coma = ";", True
    elif "," in cabecera or n_cols > 1:
        sep, coma = ",", False
    else:
        sep, coma = "\x1f", True  # una columna: ningún separador; el decimal puede ser coma
    filas = list(csv.reader(io.StringIO(texto, newline=""), delimiter=sep))
    return [list(f) for f in filas], coma


def _filas_xlsx(contenido: bytes) -> list[list[object]]:
    from openpyxl import load_workbook
    from openpyxl.xml import DEFUSEDXML

    if not DEFUSEDXML:
        raise ErrorCarga("carga.err.xml")
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_DESCOMPRIMIDO:
                raise ErrorCarga("carga.err.tamano")
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        try:
            hoja = wb.worksheets[0]
            return [
                list(f)
                for f in hoja.iter_rows(
                    max_row=MAX_FILAS + 50, max_col=MAX_COLUMNAS, values_only=True
                )
            ]
        finally:
            wb.close()
    except ErrorCarga:
        raise
    except Exception as exc:  # noqa: BLE001 - cualquier fallo de un archivo ajeno es «ilegible»
        raise ErrorCarga("carga.err.ilegible") from exc


# ------------------------------------------------------------------ la función pura
def leer_tabla(nombre: str, contenido: bytes, columnas: Sequence[Columna]) -> pd.DataFrame:
    """La tabla de un CSV o XLSX con las `columnas` pedidas (en ese orden, como float; una columna
    opcional ausente queda en NaN). Lanza `ErrorCarga` si el archivo no sirve."""
    ext = PurePath(str(nombre)).suffix.lower()
    if ext not in EXTENSIONES:
        raise ErrorCarga("carga.err.ext")
    if len(contenido) > MAX_BYTES:
        raise ErrorCarga("carga.err.tamano")
    if ext == ".csv":
        filas, coma = _filas_csv(contenido, len(columnas))
    else:
        filas, coma = _filas_xlsx(contenido), True
    # (número de fila en la hoja, celdas): las filas totalmente vacías se saltan
    con_datos = [(n, f) for n, f in enumerate(filas, start=1) if not all(_vacia(c) for c in f)]
    if not con_datos:
        raise ErrorCarga("carga.err.vacio")
    posicion = _posiciones(con_datos[0][1], columnas)
    datos = con_datos[1:]
    if not datos:
        raise ErrorCarga("carga.err.vacio")
    if len(datos) > MAX_FILAS:
        raise ErrorCarga("carga.err.filas", {"max": MAX_FILAS})
    cuerpo = np.full((len(datos), len(columnas)), np.nan)
    for i, (n_fila, fila) in enumerate(datos):
        for j, col in enumerate(columnas):
            if j not in posicion:
                continue
            celda = fila[posicion[j]] if posicion[j] < len(fila) else None
            if _vacia(celda) and not col.obligatoria:
                continue
            v = _numero(celda, coma)
            if v is None:
                raise ErrorCarga("carga.err.celda", {"fila": n_fila, "col": col})
            cuerpo[i, j] = v
    return pd.DataFrame(cuerpo, columns=[c.clave for c in columnas])


def _posiciones(cabecera: list[object], columnas: Sequence[Columna]) -> dict[int, int]:
    """Índice de columna del archivo para cada columna pedida; lanza si falta una obligatoria."""
    nombres = [_normal(str(c)) if c is not None else "" for c in cabecera]
    pos: dict[int, int] = {}
    for j, col in enumerate(columnas):
        for k, n in enumerate(nombres):
            if n in col.alias and k not in pos.values():
                pos[j] = k
                break
    faltan = [c for j, c in enumerate(columnas) if j not in pos and c.obligatoria]
    if faltan:
        raise ErrorCarga("carga.err.faltan", {"cols": faltan})
    return pos


# ------------------------------------------------------------------ la plantilla
def plantilla_csv(columnas: Sequence[Columna], ejemplo: pd.DataFrame, lang: str) -> bytes:
    """La plantilla: encabezados del idioma y las filas de ejemplo, en el formato de Excel de ese
    idioma (español: «;» y coma decimal; inglés: «,» y punto). UTF-8 con BOM para que Excel la abra
    bien. Sale de las mismas `Columna` que leen la tabla."""
    sep = ";" if lang != "en" else ","
    salida = io.StringIO(newline="")
    w = csv.writer(salida, delimiter=sep, lineterminator="\r\n")
    w.writerow([c.titulo(lang) for c in columnas])
    for _, fila in ejemplo.iterrows():
        celdas = []
        for c in columnas:
            v = fila.get(c.clave, np.nan)
            texto = "" if pd.isna(v) else f"{v:g}"
            celdas.append(texto.replace(".", ",") if sep == ";" else texto)
        w.writerow(celdas)
    return ("﻿" + salida.getvalue()).encode("utf-8")


# ------------------------------------------------------------------ la interfaz
def mensaje_error(e: ErrorCarga) -> str:
    """El texto traducido de un rechazo, con las columnas escritas en el idioma de la app."""
    from i18n import current, t

    lang = current()
    campos: dict[str, object] = {}
    for k, v in e.campos.items():
        if isinstance(v, Columna):
            campos[k] = v.titulo(lang)
        elif isinstance(v, list):
            campos[k] = ", ".join(c.titulo(lang) for c in v)
        else:
            campos[k] = v
    return t(e.clave, **campos)


def cargador(
    prefijo: str, columnas: Sequence[Columna], ejemplo: pd.DataFrame
) -> pd.DataFrame | None:
    """Dibuja el aviso, el cargador y el botón de la plantilla. Devuelve la tabla leída solo en la
    corrida en que se subió un archivo nuevo (si no, None); un rechazo se muestra y devuelve None.
    El estado del cargador vive bajo `carga_<prefijo>_`, fuera del prefijo de la pestaña, para que
    «volver al ejemplo» no vuelva a leer el archivo que sigue en el cargador."""
    import streamlit as st
    from i18n import current, t

    st.caption(t("carga.aviso"))
    subido = st.file_uploader(
        t("carga.label"), type=["csv", "xlsx"], key=f"carga_{prefijo}_archivo", help=t("carga.help")
    )
    st.download_button(
        t("carga.plantilla"),
        plantilla_csv(columnas, ejemplo, current()),
        f"plantilla_{prefijo}.csv",
        "text/csv",
        key=f"carga_{prefijo}_plantilla",
    )
    if subido is None:
        st.session_state.pop(f"carga_{prefijo}_id", None)
        return None
    marca = (subido.file_id, subido.size)
    if st.session_state.get(f"carga_{prefijo}_id") == marca:
        return None
    st.session_state[f"carga_{prefijo}_id"] = marca
    try:
        tabla = leer_tabla(subido.name, subido.getvalue(), columnas)
    except ErrorCarga as e:
        st.error(mensaje_error(e))
        return None
    st.success(t("carga.ok", n=len(tabla)))
    return tabla
